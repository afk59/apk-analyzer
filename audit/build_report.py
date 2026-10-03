# -*- coding: utf-8 -*-
"""Builds AUDIT_REPORT.md and AUDIT_REPORT.xlsx (repo root) from report_content.py
plus the measured outputs in audit/out/. Run: python audit/build_report.py"""
import json
import sys
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "audit" / "out"
sys.path.insert(0, str(Path(__file__).parent))
import report_content as C  # noqa: E402


def load(name):
    p = OUT / name
    return json.loads(p.read_text(encoding="utf-8")) if p.exists() else None


art_corpus = load("artifacts_check_corpus_run.json")
art_ida = load("artifacts_check_h_ida.json")
reg = load("regression_check.json")
corpus = load("corpus_summary.json") or {}


def art_text():
    parts = []
    for name, a in (("IDA (52 EXE/DLL)", art_ida), ("Korpus (APK)", art_corpus)):
        if a:
            real = [x for x in a["problems"] if x["app"] != "notepad.exe"]
            extra = (" (notepad.exe satırı bu denetimin paralel koşu testinden — BUG-03 — kaldı, korpus batch'inin parçası değil)"
                     if len(real) != a["problem_count"] else "")
            parts.append(f"{name}: {a['apps_checked'] - (a['problem_count'] - len(real))} uygulama, {len(real)} tutarsızlık{extra}")
        else:
            parts.append(f"{name}: koşturulamadı")
    return "; ".join(parts) + "."


def reg_text():
    if not reg:
        return "Korpus koşusu tamamlanmadığı için koşturulamadı."
    gained = Counter(k for r in reg["rows"] for k in r.get("added", []))
    top = ", ".join(f"{k} ({v})" for k, v in gained.most_common(6))
    return (f"{reg['compared']} uygulama karşılaştırıldı; kit kaybeden uygulama: {reg['apps_losing_a_kit']} "
            f"(toplam {reg['kits_lost_total']} kit). Kazanılan kit'ler (uygulama sayısı): {top}.")


def durum(res):
    r = res.lower()
    if "HATALI" in res or "güvenlik" in r:
        return "Hatalı"
    if r.startswith("kısmen"):
        return "Kısmen / hatalı"
    if "kod incelemesi" in r or "koşturulmadı" in r or "edilmedi" in r:
        return "Doğrulanmadı (yalnız kod/veri)"
    if r in ("çalışıyor (doğrulandı)", "örneklemde doğru"):
        return "Çalışıyor (canlı doğrulandı)"
    return "Çalışıyor, kusurlu"


FEATURES = [f[:8] + (durum(f[7]),) + f[8:] for f in C.FEATURES]

subs = {"%%ARTIFACTS%%": art_text(), "%%REGRESSION%%": reg_text()}
CROSS = [tuple(subs.get(c, c) if isinstance(c, str) else c for c in row) for row in C.CROSS]
CROSS = [tuple(next((v for k, v in subs.items() if k in c), c) if isinstance(c, str) else c for c in row)
         for row in CROSS]

# ------------------------------------------------------------------ markdown
sev_order = {"Kritik": 0, "Orta": 1, "Düşük": 2}
cats = [C.B, C.F, C.G]
cnt = Counter((f[1], f[2]) for f in C.FINDINGS)


def md_table(headers, rows):
    esc = lambda s: str(s).replace("|", "\\|").replace("\n", " ")  # noqa: E731
    out = ["| " + " | ".join(headers) + " |", "|" + "---|" * len(headers)]
    out += ["| " + " | ".join(esc(c) for c in r) + " |" for r in rows]
    return "\n".join(out)


md = []
md.append("# AUDIT_REPORT — Migration Analyzer v2 (apk-analyzer-v2, branch `v2`)\n")
md.append("Denetim tarihi: 2026-10-03 (gece, gözetimsiz). Denetlenen commit: `700cac7f` "
          "(çalışma ağacında yalnız `play_metadata_cache.json` değişik). Excel karşılığı: `AUDIT_REPORT.xlsx` "
          "(sheet'ler: Envanter, Feature Denetimi, Çapraz Doğrulama, Bulgular, Özet).\n")
md.append("> **Doğrulama etiketleri** — **CANLI**: diskteki gerçek dosyalarla koşturuldu ve çıktı incelendi · "
          "**SCRIPT**: `audit/` altındaki tarama script'iyle ölçüldü · **KOD**: yalnız kod incelemesi, "
          "koşturulmadı · **DOK**: doküman iddiası karşılaştırması. Test edemediğim hiçbir şey 'çalışıyor' "
          "olarak yazılmadı.\n")
md.append("## Adım 1 — Envanter (uygulamanın gerçek mevcut durumu)\n")
md.append(md_table(["ID", "Katman", "Bileşen / feature", "Tür", "Dosya(lar)", "Ne yapıyor", "Giriş noktası"], C.INVENTORY))
md.append("\n## Yönetici özeti\n")
crit = [f for f in C.FINDINGS if f[2] == "Kritik"]
md.append(f"- Toplam **{len(C.FINDINGS)} bulgu**: "
          + ", ".join(f"{c}: {sum(v for (cc, s), v in cnt.items() if cc == c)}" for c in cats)
          + f". Önem: Kritik {sum(1 for f in C.FINDINGS if f[2]=='Kritik')}, "
            f"Orta {sum(1 for f in C.FINDINGS if f[2]=='Orta')}, Düşük {sum(1 for f in C.FINDINGS if f[2]=='Düşük')}.")
md.append("- **Kritik:** " + "; ".join(f"{f[0]} — {f[3]}" for f in crit) + ".")
md.append("- Uygulama kabuğu (başlatma, API korumaları, SSE, indirici, iptal, IDA 52-binary batch, IPA, combined mod) "
          "gerçek dosyalarla **çalışıyor**; mevcut 258 test yeşil. Sorunların ağırlığı **mapping verisi/eşleme mantığında** "
          "(GMS satırları ölü, HMS ve Windows'ta yanlış eşlemeler, izin hedefleri) ve **feature'lar arası etkileşimde** "
          "(paralel koşu, taşınabilirlik, PC toplu kitabı, rapor güvenliği).")
md.append("- Korpus (42 APK) bu denetimde güncel mapping v1.14.0 ile uygulamanın kendi API'si üzerinden izole bir klasörde "
          f"yeniden koşturuldu: `audit/out/corpus_run/result`. {corpus.get('text', '')}")
md.append("\n## Adım 2 — Feature denetimi\n")
md.append(md_table(["ID", "Feature", "İddia", "Kod bunu yapıyor mu", "Edge-case / hata yönetimi", "Doğrulama",
                    "Test girdisi", "Sonuç", "Durum", "Bulgular"], FEATURES))
md.append("\n## Adım 3 — Çapraz kesen doğrulamalar\n")
md.append(md_table(["ID", "Alan", "Kontrol", "Yöntem / script", "Kapsam", "Sonuç", "Bulgu"], CROSS))
md.append("\n## Adım 4 — Rapor: bulgular\n")
for cat in cats:
    md.append(f"\n### {cat}\n")
    rows = sorted([f for f in C.FINDINGS if f[1] == cat], key=lambda f: (sev_order[f[2]], f[0]))
    for f in rows:
        md.append(f"#### {f[0]} · **{f[2]}** · {f[3]}\n")
        md.append(f"- **Açıklama:** {f[4]}")
        md.append(f"- **Kanıt:** {f[5]}")
        md.append(f"- **Doğrulama:** {f[6]}")
        md.append(f"- **Öneri:** {f[7]}\n")
md.append("\n## Neyi koşturdum, neyi koşturmadım\n")
md.append("""**Gerçekten koşturulanlar (CANLI):**
- Mevcut test paketi: `MA_TEST_APK=com.mtr.mobile.apk python -m unittest discover -s tests -t .` → 258 test, 256 sn, OK.
- Uygulama sunucusu `python -m app --no-window` (izole `MA_RESULT_DIR/MA_DB_PATH/MA_INPUT_DIR`), API üzerinden 42 APK'lık korpus koşusu (mapping v1.14.0), paralel ikinci koşu (notepad.exe), SSE/Last-Event-ID, 21 negatif HTTP isteği.
- Worker: com.mtr.mobile.apk (hmos + hms), 3 sentetik IPA (geçerli / bozuk zip / Info.plist'siz), 4 enjeksiyon IPA'sı (headless Edge ile yürütme kanıtı).
- JobRunner harness: IDA Free 9.3 program klasörü (52 binary, alt klasörler, ida.exe/ida.dll çakışması), Cppcheck combined, iptal testi (iLearn + calculator).
- v1 CLI `--batch` izole motor kopyasında (4 EXE/DLL).
- İndirici: APKPure'den gerçek indirme (com.google.android.calculator, 6,2 MB), olmayan paket, 'C:evil'.
- Kütüphane: rescan değişmezi, silme, gerçek klasör taşıma.
- `exe_engine._match_kits` System32'deki 3.558 DLL adıyla.

**Yalnız kod incelemesiyle değerlendirilenler (KOD):** yerel dosya diyaloğu (gece GUI etkileşimi yapılmadı), frontend buton durumları (paralellik API düzeyinde canlı gösterildi), redrawStages, stderr kilitlenme riski, EXE/IPA güven politikası, indirme+analiz örtüşmesi (mevcut testler kapsıyor), OHPM sweep aracı ve API index üreticisi (ağ/uzun süre), APK yolu için rapor enjeksiyonu (aynı sink; ikili manifest üretmek gerekirdi).

**Test edilemeyenler:** Duolingo'nun androguard çökmesi (APK artık diskte yok), HSBC/Calm (bakımcı kararıyla kaldırılmış), şirket ağı/XGATE davranışı, başka bir makinede ilk kurulum.

**Denetimin bıraktığı dosyalar (hepsi `audit/` altında, commit edilmedi):** `scan_mappings.py`, `check_permissions.py`, `check_i18n.py`, `check_matching.py`, `check_library.py`, `check_artifacts.py`, `check_regression.py`, `harness_runner.py`, `report_content.py`, `build_report.py`; ölçüm çıktıları `audit/out/*.json`; korpus çıktıları `audit/out/corpus_run/`. Gerçek `engine/result/` ve `app/store/analyzer.db`'ye dokunulmadı.
""")
(ROOT / "AUDIT_REPORT.md").write_text("\n".join(md), encoding="utf-8")

# ------------------------------------------------------------------ xlsx
FONT = "Arial"
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
SEV_FILL = {"Kritik": "F8CBAD", "Orta": "FFE699", "Düşük": "E2EFDA"}

wb = Workbook()


def sheet(title, headers, rows, widths, note=None, table=True, sev_col=None):
    ws = wb.create_sheet(title)
    r0 = 1
    if note:
        ws.cell(1, 1, note).font = Font(name=FONT, italic=True, size=9, color="595959")
        r0 = 3
    for j, h in enumerate(headers, 1):
        c = ws.cell(r0, j, h)
        c.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        c.fill = HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = BORDER
    for i, row in enumerate(rows, r0 + 1):
        for j, v in enumerate(row, 1):
            c = ws.cell(i, j, v)
            c.font = Font(name=FONT, size=9)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.border = BORDER
        if sev_col is not None and row[sev_col] in SEV_FILL:
            ws.cell(i, sev_col + 1).fill = PatternFill("solid", fgColor=SEV_FILL[row[sev_col]])
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = ws.cell(r0 + 1, 2)
    if table and rows:
        ref = f"A{r0}:{get_column_letter(len(headers))}{r0 + len(rows)}"
        t = Table(displayName=title.replace(" ", "_").replace("ç", "c").replace("ğ", "g").replace("Ç", "C")
                  .replace("ı", "i").replace("ş", "s").replace("ö", "o").replace("ü", "u"), ref=ref)
        t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
        ws.add_table(t)
    return ws


NOTE = ("Doğrulama: CANLI = gerçek dosyalarla koşturuldu · SCRIPT = audit/ script'i ile ölçüldü · "
        "KOD = yalnız kod incelemesi · DOK = doküman karşılaştırması. Ayrıntı: AUDIT_REPORT.md")
wb.remove(wb.active)
sheet("Envanter", ["ID", "Katman", "Bileşen / feature", "Tür", "Dosya(lar)", "Ne yapıyor / iddia", "Giriş noktası"],
      C.INVENTORY, [8, 13, 34, 14, 38, 70, 34], note="Adım 1 — uygulamanın gerçek mevcut durumu (kod okunarak çıkarıldı).")
sheet("Feature Denetimi", ["ID", "Feature", "İddia", "Kod bunu yapıyor mu", "Edge-case / hata yönetimi", "Doğrulama",
                           "Test girdisi", "Sonuç", "Durum", "Bulgular"],
      FEATURES, [7, 28, 34, 50, 50, 26, 30, 24, 18, 16], note="Adım 2 — " + NOTE)
sheet("Çapraz Doğrulama", ["ID", "Alan", "Kontrol", "Yöntem / script", "Kapsam", "Sonuç", "Bulgu"],
      CROSS, [7, 26, 44, 30, 24, 70, 14], note="Adım 3 — " + NOTE)
FROWS = sorted(C.FINDINGS, key=lambda f: (cats.index(f[1]), sev_order[f[2]], f[0]))
ws = sheet("Bulgular", ["ID", "Kategori", "Önem", "Başlık", "Açıklama", "Kanıt (dosya:satır)", "Doğrulama", "Öneri"],
           FROWS, [9, 18, 9, 40, 70, 50, 16, 50], note="Adım 4 — " + NOTE, sev_col=2)
n = len(FROWS)
first, last = 4, 3 + n

# Özet: formulas over the Bulgular sheet
s = wb.create_sheet("Özet", 0)
s["A1"] = "Denetim özeti — Migration Analyzer v2"
s["A1"].font = Font(name=FONT, bold=True, size=13)
s["A2"] = "Sayılar 'Bulgular' sheet'inden formülle hesaplanır."
s["A2"].font = Font(name=FONT, italic=True, size=9, color="595959")
hdr = ["Kategori", "Kritik", "Orta", "Düşük", "Toplam"]
for j, h in enumerate(hdr, 1):
    c = s.cell(4, j, h)
    c.font = Font(name=FONT, bold=True, color="FFFFFF")
    c.fill = HEAD_FILL
    c.border = BORDER
for i, cat in enumerate(cats, 5):
    s.cell(i, 1, cat).font = Font(name=FONT)
    for j, sev in enumerate(["Kritik", "Orta", "Düşük"], 2):
        s.cell(i, j, f"=COUNTIFS(Bulgular!$B${first}:$B${last},$A{i},Bulgular!$C${first}:$C${last},{get_column_letter(j)}$4)")
    s.cell(i, 5, f"=SUM(B{i}:D{i})")
tr = 5 + len(cats)
s.cell(tr, 1, "Toplam").font = Font(name=FONT, bold=True)
for j in range(2, 6):
    L = get_column_letter(j)
    s.cell(tr, j, f"=SUM({L}5:{L}{tr - 1})").font = Font(name=FONT, bold=True)
for r in range(5, tr + 1):
    for j in range(1, 6):
        s.cell(r, j).border = BORDER
        if j > 1:
            s.cell(r, j).font = Font(name=FONT, bold=(r == tr))
s.cell(tr + 2, 1, "Feature durumları (Feature Denetimi!I)").font = Font(name=FONT, bold=True)
labels = ["Çalışıyor (canlı doğrulandı)", "Çalışıyor, kusurlu", "Kısmen / hatalı", "Hatalı",
          "Doğrulanmadı (yalnız kod/veri)"]
ffirst, flast = 4, 3 + len(C.FEATURES)
s.cell(tr + 3 + len(labels), 1, "Toplam feature").font = Font(name=FONT, size=9, bold=True)
s.cell(tr + 3 + len(labels), 2, f"=SUM(B{tr + 3}:B{tr + 2 + len(labels)})").font = Font(name=FONT, size=9, bold=True)
for k, lab in enumerate(labels):
    r = tr + 3 + k
    s.cell(r, 1, lab).font = Font(name=FONT, size=9)
    s.cell(r, 2, f"=COUNTIF('Feature Denetimi'!$I${ffirst}:$I${flast},A{r})").font = Font(name=FONT, size=9)
s.column_dimensions["A"].width = 52
for L in "BCDE":
    s.column_dimensions[L].width = 11
wb.save(ROOT / "AUDIT_REPORT.xlsx")
print("wrote AUDIT_REPORT.md and AUDIT_REPORT.xlsx;", len(C.FINDINGS), "findings")
