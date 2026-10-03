"""Audit script 6 -- cross-artifact consistency for one finished run directory.

For every analysed input listed in the run manifest(s), compares the numbers
each artifact states about it:

  HTML report_data  <->  per-app XLSX Overview  <->  agent JSON  <->  side-car
  <->  manifest row  <->  SQLite index row  <->  batch_summary.xlsx row

    python audit/check_artifacts.py <run_out_dir>     (holds result/ and analyzer.db)
"""
import json
import re
import sqlite3
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "engine"))
from core.renderer import coverage_status  # noqa: E402
import openpyxl  # noqa: E402

base = Path(sys.argv[1]).resolve()
res = base / "result"
RD = re.compile(r"const REPORT_DATA = (\{.*?\});\n", re.S)

batch = {}
bs = res / "batch_summary.xlsx"
if bs.exists():
    wb = openpyxl.load_workbook(bs, read_only=True)
    rows = list(wb["Batch Summary"].iter_rows(values_only=True))
    hdr = rows[1]
    for r in rows[2:]:
        if r and r[0]:
            batch[r[0]] = dict(zip(hdr, r))
    wb.close()

db = sqlite3.connect(base / "analyzer.db")
db.row_factory = sqlite3.Row
idx = {}
for a in db.execute("SELECT a.*, (SELECT path FROM artifacts WHERE app_id=a.id AND kind='html') AS html FROM apps a"):
    if a["html"]:
        idx.setdefault(Path(a["html"]).name, []).append(dict(a))

problems, checked = [], 0
for mf in sorted((res / ".runs").glob("run-*.json")):
    man = json.loads(mf.read_text(encoding="utf-8"))
    for row in man.get("results", []):
        if row.get("skipped") or not row.get("report"):
            continue
        checked += 1
        name = row["source_file"]
        rep = Path(row["report"])
        P = lambda what, a, b: problems.append({"app": name, "check": what, "a": a, "b": b})  # noqa: E731
        html = rep.read_text(encoding="utf-8")
        m = RD.search(html)
        rd = json.loads(m.group(1)) if m else {}
        agent = json.loads(Path(row["agent_json"]).read_text(encoding="utf-8"))
        side = json.loads(Path(row["sidecar"]).read_text(encoding="utf-8")) if row.get("sidecar") else {}
        mode = row.get("mode")
        kits_rd = len(rd.get("kit_matches", []))
        if kits_rd != row["kit_count"]:
            P("kit_count report_data vs manifest", kits_rd, row["kit_count"])
        if len(agent.get("kit_list", [])) != row["kit_count"]:
            P("kit_count agent vs manifest", len(agent.get("kit_list", [])), row["kit_count"])
        if mode == "apk":
            un_rd = len(rd.get("unmapped_namespaces", []))
            if un_rd != row["unmapped_count"]:
                P("unmapped report_data vs manifest", un_rd, row["unmapped_count"])
            if len(side.get("unmapped", [])) != row["unmapped_count"]:
                P("unmapped sidecar vs manifest", len(side.get("unmapped", [])), row["unmapped_count"])
            for k in ("total_detected_ns", "total_3p_ns", "total_ohpm_ns"):
                if rd.get(k) != side.get(k):
                    P(k + " report_data vs sidecar", rd.get(k), side.get(k))
            pct, st = coverage_status({**rd, "framework": rd.get("framework") or {}})
            if (pct, st) != (row.get("coverage_pct"), row.get("coverage_status")):
                P("coverage recomputed vs manifest", (pct, st), (row.get("coverage_pct"), row.get("coverage_status")))
            # the HTML stat card shows the same number?
            if st == "ok" and f"{pct}%" not in html:
                P("coverage % not visible in HTML", pct, None)
            if row.get("xlsx"):
                wb = openpyxl.load_workbook(row["xlsx"], read_only=True)
                ov = {r[0]: r[1] for r in wb["Overview"].iter_rows(values_only=True) if r and r[0]}
                wb.close()
                if ov.get("Kit Matches") != row["kit_count"]:
                    P("kit_count xlsx vs manifest", ov.get("Kit Matches"), row["kit_count"])
                if ov.get("3rd Party (OHPM)") != row.get("ohpm_count"):
                    P("ohpm xlsx vs manifest", ov.get("3rd Party (OHPM)"), row.get("ohpm_count"))
        ix = idx.get(rep.name) or []
        if len(ix) != 1:
            P("index rows for this report", len(ix), 1)
        else:
            i = ix[0]
            for k in ("kit_count", "unmapped_count", "coverage_pct", "coverage_status", "ohpm_count"):
                if i[k] != row.get(k):
                    P(k + " index vs manifest", i[k], row.get(k))
        b = batch.get(name)
        if batch and b is None:
            P("row missing from batch_summary", None, name)
        elif b:
            if b.get("Kit Count") != row["kit_count"]:
                P("kit_count batch vs manifest", b.get("Kit Count"), row["kit_count"])
            if b.get("Unmapped") != row["unmapped_count"]:
                P("unmapped batch vs manifest", b.get("Unmapped"), row["unmapped_count"])
            if mode == "apk":
                cov = b.get("Coverage %")
                exp = row.get("coverage_pct")
                if exp is not None and (not isinstance(cov, (int, float)) or abs(cov - exp) > 0.051):
                    P("coverage batch vs manifest", cov, exp)

out = {"run_dir": str(base), "apps_checked": checked, "batch_rows": len(batch),
       "problems": problems, "problem_count": len(problems)}
(ROOT / "audit" / "out" / f"artifacts_check_{base.name}.json").write_text(
    json.dumps(out, indent=1, ensure_ascii=False, default=str), encoding="utf-8")
print(f"{base.name}: apps checked {checked}, batch rows {len(batch)}, problems {len(problems)}")
for p in problems[:40]:
    print("  ", p)
