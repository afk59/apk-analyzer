"""Audit harness -- drive the REAL JobRunner (worker subprocesses, aggregate,
manifest, SQLite rescan) over an input, in an isolated result directory, and
check the invariants the docs promise.

    python audit/harness_runner.py <input> <outdir> [--pc-mode batch|combined]
                                   [--no-recursive] [--resources] [--target hmos|hms]
                                   [--skip]

Same code path as POST /api/runs minus HTTP: input expansion comes from
app.api.routes._inputs_for, the queue from app.services.jobs.JobRunner.
"""
import argparse
import json
import os
import sys
import time
from pathlib import Path

ap = argparse.ArgumentParser()
ap.add_argument("input")
ap.add_argument("outdir")
ap.add_argument("--pc-mode", default="batch")
ap.add_argument("--no-recursive", action="store_true")
ap.add_argument("--resources", action="store_true")
ap.add_argument("--target", default="hmos")
ap.add_argument("--skip", action="store_true")
a = ap.parse_args()

out = Path(a.outdir).resolve()
out.mkdir(parents=True, exist_ok=True)
os.environ["MA_RESULT_DIR"] = str(out / "result")
os.environ["MA_DB_PATH"] = str(out / "analyzer.db")
os.environ["MA_INPUT_DIR"] = str(out / "input")
ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from app import config  # noqa: E402
config.ensure_dirs()
from app.api.routes import _inputs_for  # noqa: E402
from app.services.jobs import JobRunner  # noqa: E402
from app.store import index  # noqa: E402

inp = Path(a.input)
files, label, root = _inputs_for(inp, a.pc_mode, None if not a.no_recursive else False, a.resources)
skip = a.skip
r = JobRunner()
t0 = time.time()
rid = r.submit(files, a.target, skip_existing=skip, input_label=label, root=root)
run = r.get(rid)
cursor = 0
while True:
    evs, cursor, fin = r.events(rid, cursor, timeout=5)
    if fin and not evs:
        break
snap = run.snapshot()
elapsed = round(time.time() - t0, 1)

checks = {}
reports = [x.get("report") for x in snap["results"] if x.get("report")]
checks["files_planned"] = len(files)
checks["results"] = len(snap["results"])
checks["failures"] = len(snap["failures"])
checks["skipped"] = sum(1 for x in snap["results"] if x.get("skipped"))
checks["distinct_report_paths"] = len(set(map(str.lower, reports)))
checks["report_paths_all_distinct"] = len(set(map(str.lower, reports))) == len(reports)
checks["reports_exist_on_disk"] = all(Path(p).exists() for p in reports)
checks["stems_overridden"] = run.stems
checks["workbooks"] = snap["workbooks"]
lib = index.library()
checks["index_runs"] = len(lib["runs"])
checks["index_apps"] = sum(len(x["apps"]) for x in lib["runs"])
checks["index_failed_rows"] = sum(1 for x in lib["runs"] for ap_ in x["apps"] if ap_["status"] == "failed")
manifests = sorted((out / "result" / ".runs").glob("run-*.json"))
checks["manifests"] = [m.name for m in manifests]
if snap["workbooks"]:
    import openpyxl
    wb = openpyxl.load_workbook(snap["workbooks"]["batch_summary"], read_only=True)
    checks["batch_summary_sheets"] = wb.sheetnames
    ws = wb[wb.sheetnames[0]]
    rows = [r_ for r_ in ws.iter_rows(values_only=True)]
    checks["batch_summary_first_sheet_rows"] = len(rows)
    checks["batch_summary_header"] = rows[0] if rows else None
    wb.close()
summary = {"label": label, "elapsed_sec": elapsed, "status": snap["status"],
           "checks": checks,
           "failures_detail": snap["failures"][:20],
           "events": len(run.events)}
(out / "harness_summary.json").write_text(json.dumps(summary, indent=1, default=str, ensure_ascii=False),
                                          encoding="utf-8")
(out / "events.jsonl").write_text("\n".join(json.dumps(e, default=str, ensure_ascii=False) for e in run.events),
                                  encoding="utf-8")
print(json.dumps(summary, indent=1, default=str, ensure_ascii=False))
