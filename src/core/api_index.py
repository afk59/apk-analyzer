import re
import json
from pathlib import Path

_XLSX_PATH  = Path(__file__).parent.parent / "data" / "reference" / "hmos_api.xlsx"
_INDEX_PATH = Path(__file__).parent.parent / "data" / "reference" / "hmos_api_index.json"
_VER_RE     = re.compile(r'\d+\+?\s*$')


def _norm(s: str) -> str:
    return _VER_RE.sub('', s).strip().lower()


def _simple(name: str) -> str:
    return _norm(name.split('.')[-1])


def build_index() -> dict:
    import openpyxl
    wb = openpyxl.load_workbook(_XLSX_PATH, read_only=True, data_only=True)
    ws = wb["API Device Matrix"]
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    col = {h: i for i, h in enumerate(headers)}

    index = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        raw = row[col["API / Method"]]
        if not raw:
            continue
        api_method = str(raw).strip()
        key = _simple(api_method)
        if key not in index:
            index[key] = {
                "kit":              row[col.get("Kit")],
                "import_statement": row[col.get("Import Statement")],
                "doc_url":          row[col.get("URL")],
            }
    wb.close()
    return index


if __name__ == "__main__":
    print(f"Reading {_XLSX_PATH} ...")
    index = build_index()
    _INDEX_PATH.write_text(
        json.dumps(index, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    print(f"Wrote {len(index)} entries to {_INDEX_PATH}")
